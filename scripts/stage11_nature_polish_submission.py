from __future__ import annotations

import math
import re
import shutil
import textwrap
from pathlib import Path
from typing import Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from sklearn.calibration import calibration_curve
from sklearn.metrics import average_precision_score, precision_recall_curve, roc_auc_score, roc_curve


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
SOURCE_PACKAGE = Path.home() / "Desktop" / "final_submission_package"
DEST_PACKAGE = Path.home() / "Desktop" / "final_submission_package_nature_polished"
DEST_TABLES = DEST_PACKAGE / "tables"
DEST_FIGURES = DEST_PACKAGE / "figures"
DEST_SOURCE = DEST_PACKAGE / "source_data"

TABLES = SOURCE_PACKAGE / "tables"
PROJECT_TABLES = OUT / "tables"
PROJECT_FIGURES = OUT / "figures"
EXTERNAL = OUT / "external_validation"

PALETTE = {
    "full_clinical_model": "#0F4D92",
    "routine_indicator_model": "#42949E",
    "biomarker_focused_model": "#B64342",
    "Global": "#767676",
    "China": "#0F4D92",
    "United States": "#8BCF8B",
    "reference": "#9A9A9A",
    "signal": "#0F4D92",
    "accent": "#B64342",
}

FEATURE_SET_LABELS = {
    "full_clinical_model": "Full clinical model",
    "routine_indicator_model": "Routine indicator model",
    "biomarker_focused_model": "Biomarker-focused model",
}

BEST_MODELS = {
    "full_clinical_model": "XGBoost",
    "routine_indicator_model": "XGBoost",
    "biomarker_focused_model": "Elastic Net",
}

DISPLAY = {
    "uacr": "UACR",
    "egfr": "eGFR",
    "tyg_wc": "TyG-WC",
    "siri": "SIRI",
    "tyg_index": "TyG index",
    "non_hdl_cholesterol": "non-HDL-C",
    "ldl_cholesterol": "LDL-C",
    "hba1c": "HbA1c",
    "age": "Age",
    "pir": "PIR",
    "non_hdl_cholesterol": "non-HDL-C",
    "serum_creatinine": "Serum creatinine",
    "smoking_status": "Smoking status",
    "serum_albumin": "Serum albumin",
    "nlr": "NLR",
    "plr": "PLR",
    "bmi": "BMI",
    "uric_acid": "Uric acid",
    "sii": "SII",
    "crp": "CRP",
}

TABLE1_LABEL_REPLACEMENTS = {
    "Body mass index, kg/m2": "Body mass index, kg/m2",
    "Total cholesterol, mmol/L": "Total cholesterol, mg/dL",
    "LDL-C, mmol/L": "LDL-C, mg/dL",
    "HDL-C, mmol/L": "HDL-C, mg/dL",
    "non-HDL-C, mmol/L": "non-HDL-C, mg/dL",
    "eGFR, mL/min/1.73 m2": "eGFR, mL/min/1.73 m2",
}


def ensure_dirs() -> None:
    for path in [DEST_PACKAGE, DEST_TABLES, DEST_FIGURES, DEST_SOURCE]:
        path.mkdir(parents=True, exist_ok=True)


def configure_matplotlib() -> None:
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans", "sans-serif"],
            "font.size": 7.5,
            "axes.spines.right": False,
            "axes.spines.top": False,
            "axes.linewidth": 0.8,
            "axes.labelsize": 7.5,
            "axes.titlesize": 8.2,
            "xtick.labelsize": 6.8,
            "ytick.labelsize": 6.8,
            "legend.fontsize": 6.8,
            "legend.frameon": False,
            "pdf.fonttype": 42,
            "svg.fonttype": "none",
            "savefig.facecolor": "white",
        }
    )


def panel_label(ax: plt.Axes, label: str) -> None:
    ax.text(
        -0.12,
        1.08,
        label,
        transform=ax.transAxes,
        fontsize=9.5,
        fontweight="bold",
        va="top",
        ha="left",
    )


def save_figure(fig: plt.Figure, basename: str, dpi: int = 600) -> None:
    out = DEST_FIGURES / basename
    fig.savefig(out.with_suffix(".png"), dpi=dpi, bbox_inches="tight")
    fig.savefig(out.with_suffix(".pdf"), bbox_inches="tight")
    fig.savefig(out.with_suffix(".svg"), bbox_inches="tight")
    plt.close(fig)


def fmt_p(value: object) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)) or pd.isna(value):
        return "NA"
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return "NA"
        if s.startswith("P "):
            return s.replace("P = ", "").replace("P < ", "<")
        try:
            value = float(s)
        except ValueError:
            return s
    v = float(value)
    if v < 0.001:
        return "<0.001"
    if v < 1:
        return f"{v:.3f}"
    return f"{v:.2f}"


def display_feature(name: str) -> str:
    return DISPLAY.get(str(name), str(name).replace("_", " ").title())


def publication_csv(df: pd.DataFrame, filename: str) -> None:
    df.to_csv(DEST_TABLES / filename, index=False, encoding="utf-8-sig")


def polish_tables() -> dict[str, pd.DataFrame]:
    table_outputs: dict[str, pd.DataFrame] = {}

    t1 = pd.read_csv(TABLES / "Table_1_weighted_baseline_characteristics.csv")
    t1["Characteristic"] = t1["Characteristic"].replace(TABLE1_LABEL_REPLACEMENTS)
    t1["Level/statistic"] = t1["Level/statistic"].replace({"weighted mean (SE)": "Weighted mean (SE)"})
    t1["P value"] = t1["P value"].map(fmt_p)
    t1.loc[t1.duplicated("Characteristic"), "P value"] = ""
    t1.loc[t1["Level/statistic"].eq("Weighted mean (SE)"), "Level/statistic"] = ""
    table_outputs["Table 1"] = t1
    publication_csv(t1, "Table_1_weighted_baseline_characteristics_polished.csv")

    t2 = pd.read_csv(TABLES / "Table_2_model_performance.csv")
    order = ["full_clinical_model", "routine_indicator_model", "biomarker_focused_model"]
    t2["feature set"] = pd.Categorical(t2["feature set"], categories=order, ordered=True)
    t2 = t2.sort_values("feature set").reset_index(drop=True)
    t2["Feature set"] = t2["feature set"].map(FEATURE_SET_LABELS)
    t2["Best model"] = t2["best model"]
    for col in ["AUC", "PR-AUC", "F1", "Brier score", "sensitivity", "specificity"]:
        t2[col] = pd.to_numeric(t2[col], errors="coerce").map(lambda x: f"{x:.3f}")
    t2["Delta AUC versus full clinical"] = pd.to_numeric(
        t2["AUC"], errors="coerce"
    ).sub(float(t2.loc[t2["Feature set"].eq("Full clinical model"), "AUC"].iloc[0])).map(lambda x: f"{x:+.3f}")
    t2.loc[t2["Feature set"].eq("Full clinical model"), "Delta AUC versus full clinical"] = "Reference"
    t2 = t2[
        [
            "Feature set",
            "Best model",
            "n_train",
            "n_test",
            "AUC",
            "PR-AUC",
            "F1",
            "Brier score",
            "sensitivity",
            "specificity",
            "Delta AUC versus full clinical",
        ]
    ]
    t2 = t2.rename(columns={"n_train": "Training n", "n_test": "Test n", "sensitivity": "Sensitivity", "specificity": "Specificity"})
    table_outputs["Table 2"] = t2
    publication_csv(t2, "Table_2_model_performance_polished.csv")

    t3 = pd.read_csv(TABLES / "Table_3_feature_ranking.csv")
    t3 = t3.rename(
        columns={
            "indicator": "Indicator",
            "routine top10": "Routine top 10",
            "routine consensus rank": "Routine consensus rank",
            "biomarker-focused top10": "Biomarker-focused top 10",
            "biomarker consensus rank": "Biomarker consensus rank",
            "sensitivity stability": "Sensitivity stability",
            "interpretation": "Interpretation",
        }
    )
    t3["Interpretation"] = t3["Interpretation"].str.replace("ML", "machine-learning", regex=False)
    table_outputs["Table 3"] = t3
    publication_csv(t3, "Table_3_feature_ranking_polished.csv")

    t4 = pd.read_csv(TABLES / "Table_4_weighted_regression.csv")
    t4 = t4.rename(
        columns={
            "indicator": "Indicator",
            "summary": "Interpretation",
            "Model 3 n/events": "Model 3 n/events",
            "Strict Model 4 n/events": "Strict Model 4 n/events",
        }
    )
    for col in ["Model 3 P value", "Strict Model 4 P value"]:
        t4[col] = t4[col].map(fmt_p)
    t4["Interpretation"] = t4["Interpretation"].str.replace("ML", "machine-learning", regex=False)
    table_outputs["Table 4"] = t4
    publication_csv(t4, "Table_4_weighted_regression_polished.csv")

    t5 = pd.read_csv(TABLES / "Table_5_gbd_summary.csv")
    t5 = t5.rename(
        columns={
            "location": "Location",
            "measure": "Measure",
            "1990 age-standardized rate": "1990 age-standardised rate",
            "2023 age-standardized rate": "2023 age-standardised rate",
            "percent change": "Change, 1990 to 2023",
            "trend direction": "Trend direction",
        }
    )
    t5["EAPC P value"] = t5["EAPC P value"].map(fmt_p)
    t5["Trend direction"] = t5["Trend direction"].str.capitalize()
    table_outputs["Table 5"] = t5
    publication_csv(t5, "Table_5_gbd_summary_polished.csv")

    t6 = pd.read_csv(TABLES / "Table_6_external_validation.csv")
    t6["feature_set"] = t6["feature_set"].replace(FEATURE_SET_LABELS)
    t6["indicator"] = t6["indicator"].map(lambda x: display_feature(x) if isinstance(x, str) and x else x)
    for col in ["AUC", "PR-AUC", "F1", "Brier score", "internal held-out AUC", "external minus internal AUC", "OR", "CI_lower", "CI_upper"]:
        if col in t6.columns:
            t6[col] = pd.to_numeric(t6[col], errors="coerce").map(lambda x: "" if pd.isna(x) else f"{x:.3f}")
    t6["P_value"] = t6["P_value"].map(fmt_p)
    t6 = t6.rename(
        columns={
            "analysis_type": "Analysis",
            "feature_set": "Feature set",
            "final_model_or_regression_model": "Model",
            "indicator": "Indicator",
            "n_train_2005_2018": "Training n, 2005-2018",
            "n_external_or_regression_n": "External/regression n",
            "n_external_events": "Events",
            "external minus internal AUC": "External minus internal AUC",
            "P_value": "P value",
            "p_method": "P-value method",
        }
    )
    for col in ["Training n, 2005-2018", "External/regression n", "Events"]:
        if col in t6.columns:
            t6[col] = pd.to_numeric(t6[col], errors="coerce").map(lambda x: "" if pd.isna(x) else f"{int(x)}")
    t6.loc[t6["Analysis"].eq("external_model_performance"), ["P value", "direction", "P-value method"]] = ""
    table_outputs["Table 6"] = t6
    publication_csv(t6, "Table_6_external_validation_polished.csv")

    return table_outputs


def create_excel_workbook(table_outputs: dict[str, pd.DataFrame]) -> None:
    workbook = DEST_TABLES / "publication_tables_polished.xlsx"
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        for sheet, df in table_outputs.items():
            df.to_excel(writer, sheet_name=sheet, index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(workbook)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in ws.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 11), 42)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.sheet_view.showGridLines = False
    wb.save(workbook)


def copy_source_data() -> None:
    source_files = [
        PROJECT_TABLES / "gbd_cvd_burden_trend.csv",
        PROJECT_TABLES / "gbd_risk_factor_ranking.csv",
        PROJECT_TABLES / "model_predictions_2005_2018_by_featureset.csv",
        PROJECT_TABLES / "model_performance_2005_2018_by_featureset.csv",
        PROJECT_TABLES / "feature_ranking_consensus_2005_2018_by_featureset.csv",
        PROJECT_TABLES / "subgroup_analysis.csv",
        PROJECT_TABLES / "sensitivity_analysis.csv",
        PROJECT_TABLES / "rcs_nonlinearity_tests.csv",
        EXTERNAL / "external_model_performance_2021_2023.csv",
        EXTERNAL / "external_predictions_full_clinical_model.csv",
        EXTERNAL / "external_predictions_routine_indicator_model.csv",
        EXTERNAL / "external_predictions_biomarker_focused_model.csv",
    ]
    for src in source_files:
        if src.exists():
            shutil.copy2(src, DEST_SOURCE / src.name)


def create_figure1() -> None:
    trend = pd.read_csv(PROJECT_TABLES / "gbd_cvd_burden_trend.csv")
    risk = pd.read_csv(PROJECT_TABLES / "gbd_risk_factor_ranking.csv")
    trend = trend[
        trend["location"].isin(["Global", "China", "United States"])
        & trend["sex"].eq("Both")
        & trend["age"].str.contains("Age-standardized", case=False, na=False)
        & trend["metric"].eq("Rate")
        & trend["measure"].isin(["DALYs", "Deaths", "Prevalence"])
    ].copy()

    fig = plt.figure(figsize=(7.2, 5.2), constrained_layout=True)
    gs = fig.add_gridspec(2, 3, height_ratios=[1, 1.08])
    axes = [fig.add_subplot(gs[0, i]) for i in range(3)]
    measures = [("DALYs", "Age-standardised DALY rate"), ("Deaths", "Age-standardised death rate"), ("Prevalence", "Age-standardised prevalence rate")]
    for ax, (measure, ylabel) in zip(axes, measures):
        sub = trend[trend["measure"].eq(measure)]
        for loc in ["Global", "China", "United States"]:
            s = sub[sub["location"].eq(loc)].sort_values("year")
            ax.plot(s["year"], s["value"], color=PALETTE[loc], lw=1.7, label=loc)
        ax.set_title(measure)
        ax.set_xlabel("Year")
        ax.set_ylabel(ylabel)
        ax.tick_params(length=2.5, width=0.7)
        ax.margins(x=0.02)
    panel_label(axes[0], "a")
    axes[1].legend(ncol=3, loc="upper center", bbox_to_anchor=(0.5, 1.26), handlelength=1.6)

    ax = fig.add_subplot(gs[1, :])
    risk_focus = risk[
        risk["year"].eq(2023)
        & risk["measure"].eq("DALYs")
        & risk["metric"].eq("Number")
        & risk["age"].eq("All ages")
        & risk["sex"].eq("Both")
        & risk["location"].isin(["Global", "China"])
    ].copy()
    order = (
        risk_focus[risk_focus["location"].eq("Global")]
        .sort_values("rank")["rei"]
        .head(7)
        .tolist()
    )
    risk_focus = risk_focus[risk_focus["rei"].isin(order)]
    y = np.arange(len(order))
    bar_h = 0.34
    for offset, loc in [(-bar_h / 2, "China"), (bar_h / 2, "Global")]:
        s = risk_focus[risk_focus["location"].eq(loc)].set_index("rei").reindex(order)
        vals = s["value"].to_numpy() / 1_000_000
        ax.barh(y + offset, vals, height=bar_h, color=PALETTE[loc], alpha=0.92, label=loc)
    ax.set_yticks(y)
    ax.set_yticklabels(order)
    ax.invert_yaxis()
    ax.set_xlabel("Attributable CVD DALYs in 2023, millions")
    ax.set_title("Leading attributable risk factors")
    ax.legend(loc="lower right")
    panel_label(ax, "b")
    save_figure(fig, "Figure_1_GBD_CVD_burden_and_risk_factors")


def plot_prediction_panels(pred: pd.DataFrame, basename: str, external: bool = False) -> None:
    fig, axes = plt.subplots(1, 3, figsize=(7.2, 2.35), constrained_layout=True)
    for feature_set, model in BEST_MODELS.items():
        data = pred[(pred["feature_set"].eq(feature_set)) & (pred["model"].eq(model))].copy()
        if data.empty:
            continue
        y = data["observed_cvd" if "observed_cvd" in data.columns else "cvd"].astype(int)
        p = data["predicted_probability"].astype(float)
        fpr, tpr, _ = roc_curve(y, p)
        precision, recall, _ = precision_recall_curve(y, p)
        roc_auc = roc_auc_score(y, p)
        pr_auc = average_precision_score(y, p)
        label = f"{FEATURE_SET_LABELS[feature_set]} ({model}), AUC={roc_auc:.3f}"
        axes[0].plot(fpr, tpr, lw=1.7, color=PALETTE[feature_set], label=label)
        axes[1].plot(recall, precision, lw=1.7, color=PALETTE[feature_set], label=f"{FEATURE_SET_LABELS[feature_set]}, PR-AUC={pr_auc:.3f}")
        frac_pos, mean_pred = calibration_curve(y, p, n_bins=10, strategy="quantile")
        axes[2].plot(mean_pred, frac_pos, marker="o", ms=3, lw=1.5, color=PALETTE[feature_set], label=FEATURE_SET_LABELS[feature_set])

    axes[0].plot([0, 1], [0, 1], ls="--", lw=0.9, color=PALETTE["reference"])
    axes[2].plot([0, 1], [0, 1], ls="--", lw=0.9, color=PALETTE["reference"])
    titles = ["ROC curve", "Precision-recall curve", "Calibration"]
    labels = [("False-positive rate", "True-positive rate"), ("Recall", "Precision"), ("Mean predicted probability", "Observed proportion")]
    for i, ax in enumerate(axes):
        ax.set_title(titles[i])
        ax.set_xlabel(labels[i][0])
        ax.set_ylabel(labels[i][1])
        ax.set_xlim(0, 1)
        ax.tick_params(length=2.5, width=0.7)
        panel_label(ax, chr(ord("a") + i))
    axes[0].legend(loc="lower right", fontsize=5.7)
    axes[1].legend(loc="upper right", fontsize=5.7)
    axes[2].legend(loc="upper left", fontsize=5.7)
    if external:
        axes[0].set_title("External ROC curve")
        axes[1].set_title("External precision-recall curve")
        axes[2].set_title("External calibration")
    save_figure(fig, basename)


def create_figure2() -> None:
    pred = pd.read_csv(PROJECT_TABLES / "model_predictions_2005_2018_by_featureset.csv")
    plot_prediction_panels(pred, "Figure_2_ML_model_performance")


def trim_image(img: Image.Image, threshold: int = 248) -> Image.Image:
    arr = np.asarray(img.convert("RGB"))
    mask = np.any(arr < threshold, axis=2)
    coords = np.argwhere(mask)
    if coords.size == 0:
        return img
    y0, x0 = coords.min(axis=0)
    y1, x1 = coords.max(axis=0) + 1
    pad = 20
    x0 = max(0, x0 - pad)
    y0 = max(0, y0 - pad)
    x1 = min(img.width, x1 + pad)
    y1 = min(img.height, y1 + pad)
    return img.crop((x0, y0, x1, y1))


def create_figure3() -> None:
    shap = trim_image(Image.open(PROJECT_FIGURES / "fig3_shap_summary.png"))
    ranking = pd.read_csv(PROJECT_TABLES / "feature_ranking_consensus_2005_2018_by_featureset.csv")
    ranking = ranking[ranking["feature_set"].eq("biomarker_focused_model")].sort_values("consensus_rank").head(12)
    ranking["Feature"] = ranking["feature"].map(display_feature)

    fig = plt.figure(figsize=(7.2, 5.2), constrained_layout=True)
    gs = fig.add_gridspec(1, 2, width_ratios=[1.35, 0.75])
    ax_img = fig.add_subplot(gs[0, 0])
    ax_img.imshow(shap)
    ax_img.set_axis_off()
    ax_img.set_title("Biomarker-focused SHAP summary", loc="left")
    panel_label(ax_img, "a")

    ax = fig.add_subplot(gs[0, 1])
    y = np.arange(len(ranking))[::-1]
    score = ranking["top10_method_count"].astype(float).to_numpy()
    colors = [PALETTE["signal"] if f in ["UACR", "eGFR", "TyG-WC", "SIRI"] else "#B4C0E4" for f in ranking["Feature"]]
    ax.barh(y, score, color=colors, height=0.68)
    ax.set_yticks(y)
    ax.set_yticklabels(ranking["Feature"])
    ax.set_xlabel("Top-10 method count")
    ax.set_xlim(0, max(4, score.max()) + 0.5)
    ax.set_title("Consensus feature support")
    for yi, rank, mean_rank in zip(y, ranking["consensus_rank"], ranking["mean_rank"]):
        ax.text(0.05, yi, f"rank {int(rank)}", ha="left", va="center", color="white", fontsize=6.2, fontweight="bold")
        ax.text(score.max() + 0.18, yi, f"mean rank {mean_rank:.1f}", ha="left", va="center", fontsize=6.1, color="#4D4D4D")
    panel_label(ax, "b")
    save_figure(fig, "Figure_3_SHAP_and_feature_importance")


def create_figure4() -> None:
    panels = [
        ("a", "UACR", PROJECT_FIGURES / "fig4_rcs_uacr.png"),
        ("b", "eGFR", PROJECT_FIGURES / "fig4_rcs_egfr.png"),
        ("c", "TyG-WC", PROJECT_FIGURES / "fig4_rcs_tyg_wc.png"),
        ("d", "SIRI", PROJECT_FIGURES / "fig4_rcs_siri.png"),
    ]
    fig, axes = plt.subplots(2, 2, figsize=(7.2, 5.15), constrained_layout=True)
    for ax, (lab, title, path) in zip(axes.ravel(), panels):
        img = trim_image(Image.open(path), threshold=250)
        ax.imshow(img)
        ax.set_axis_off()
        ax.set_title(title, loc="left", pad=2)
        panel_label(ax, lab)
    save_figure(fig, "Figure_4_RCS_key_indicators")


def subgroup_label(row: pd.Series) -> str:
    subgroup = str(row["subgroup"]).replace("_stage3", "").replace("_", " ")
    level = str(row["level"])
    if subgroup == "sex":
        level = {"Female": "female", "Male": "male"}.get(level, level)
    if subgroup == "diabetes":
        level = "diabetes" if level == "1" else "no diabetes"
    if subgroup == "hypertension":
        level = "hypertension" if level == "1" else "no hypertension"
    if subgroup == "obesity status":
        level = "obese" if level == "obese" else "non-obese"
    return f"{display_feature(row['indicator'])}: {subgroup}, {level}"


def create_figure5() -> None:
    subgroup = pd.read_csv(PROJECT_TABLES / "subgroup_analysis.csv")
    sensitivity = pd.read_csv(PROJECT_TABLES / "sensitivity_analysis.csv")
    core = ["uacr", "egfr", "tyg_wc", "siri"]
    sub = subgroup[subgroup["indicator"].isin(core) & subgroup["form"].eq("continuous")].copy()
    sub["Indicator"] = sub["indicator"].map(display_feature)
    sub["Label"] = sub.apply(subgroup_label, axis=1)
    sub["sort_ind"] = sub["indicator"].map({v: i for i, v in enumerate(core)})
    sub = sub.sort_values(["sort_ind", "subgroup", "level"]).reset_index(drop=True)

    fig = plt.figure(figsize=(7.2, 6.1), constrained_layout=True)
    gs = fig.add_gridspec(1, 2, width_ratios=[1.05, 0.95])
    ax = fig.add_subplot(gs[0, 0])
    y = np.arange(len(sub))[::-1]
    x = sub["OR"].astype(float).to_numpy()
    lo = sub["CI_lower"].astype(float).to_numpy()
    hi = sub["CI_upper"].astype(float).to_numpy()
    colors = sub["indicator"].map({"uacr": "#0F4D92", "egfr": "#42949E", "tyg_wc": "#9A4D8E", "siri": "#B64342"}).to_numpy()
    ax.hlines(y, lo, hi, color=colors, lw=0.9, alpha=0.85)
    ax.scatter(x, y, s=12, color=colors, zorder=3)
    ax.axvline(1, color=PALETTE["reference"], lw=0.9, ls="--")
    ax.set_yticks(y)
    ax.set_yticklabels(sub["Label"], fontsize=5.1)
    ax.set_xlabel("Odds ratio per 1-unit increase")
    ax.set_title("Subgroup estimates")
    ax.set_xscale("log")
    ax.set_xlim(0.92, 1.55)
    forest_ticks = [0.95, 1.00, 1.10, 1.20, 1.30, 1.40, 1.50]
    ax.set_xticks(forest_ticks)
    ax.set_xticklabels([f"{tick:.2f}" if tick < 1 else f"{tick:.1f}" for tick in forest_ticks])
    ax.tick_params(length=2.5, width=0.7)
    panel_label(ax, "a")

    ax2 = fig.add_subplot(gs[0, 1])
    sens_order = [
        ("exclude_prior_stroke", "Exclude stroke"),
        ("exclude_prior_heart_attack", "Exclude heart attack"),
        ("routine_biomarker_adjustment_no_htn_dm", "Routine biomarkers"),
    ]
    mat = np.zeros((len(core), len(sens_order)))
    labels = [["" for _ in sens_order] for _ in core]
    for i, ind in enumerate(core):
        for j, (sens_name, _) in enumerate(sens_order):
            row = sensitivity[(sensitivity["indicator"].eq(ind)) & (sensitivity["sensitivity"].eq(sens_name))]
            if row.empty:
                mat[i, j] = np.nan
                labels[i][j] = "NA"
            else:
                r = row.iloc[0]
                p = float(r["P_value"])
                mat[i, j] = 1 if p < 0.05 else 0
                labels[i][j] = f"OR {float(r['OR']):.3f}\nP {fmt_p(p)}"
    cmap = matplotlib.colors.ListedColormap(["#F0F0F0", "#0F4D92"])
    ax2.imshow(np.nan_to_num(mat), vmin=0, vmax=1, cmap=cmap, aspect="auto")
    ax2.set_yticks(np.arange(len(core)))
    ax2.set_yticklabels([display_feature(x) for x in core])
    ax2.set_xticks(np.arange(len(sens_order)))
    ax2.set_xticklabels([x[1] for x in sens_order], rotation=25, ha="right")
    for i in range(mat.shape[0]):
        for j in range(mat.shape[1]):
            color = "white" if mat[i, j] == 1 else "#272727"
            ax2.text(j, i, labels[i][j], ha="center", va="center", fontsize=6.3, color=color)
    ax2.set_title("Sensitivity analysis")
    ax2.tick_params(length=0)
    for spine in ax2.spines.values():
        spine.set_visible(False)
    panel_label(ax2, "b")
    save_figure(fig, "Figure_5_Subgroup_and_sensitivity_analysis")


def create_figure6() -> None:
    frames = []
    for path in [
        EXTERNAL / "external_predictions_full_clinical_model.csv",
        EXTERNAL / "external_predictions_routine_indicator_model.csv",
        EXTERNAL / "external_predictions_biomarker_focused_model.csv",
    ]:
        frames.append(pd.read_csv(path))
    pred = pd.concat(frames, ignore_index=True)
    pred = pred.rename(columns={"cvd": "observed_cvd"})
    plot_prediction_panels(pred, "Figure_6_External_validation_performance", external=True)


def create_figures() -> None:
    create_figure1()
    create_figure2()
    create_figure3()
    create_figure4()
    create_figure5()
    create_figure6()


def build_legends() -> str:
    return """# Table and Figure Legends

## Tables

**Table 1. Survey-weighted baseline characteristics according to cardiovascular disease status.** Continuous variables are presented as survey-weighted means with standard errors. Categorical variables are presented as weighted percentages. Cardiovascular disease (CVD) was defined as self-reported congestive heart failure, coronary heart disease, angina, heart attack or myocardial infarction, or stroke.

**Table 2. Held-out performance of leakage-audited machine-learning models in NHANES 2005-2018.** The table shows the highest-performing model within each predefined feature set. AUC, PR-AUC, F1 score, Brier score, sensitivity and specificity were evaluated in the held-out test set.

**Table 3. Feature-prioritisation and statistical validation summary for selected indicators.** The table integrates consensus feature rankings, survey-weighted regression, restricted cubic spline testing and sensitivity analyses for priority renal, metabolic, inflammatory and lipid-related indicators.

**Table 4. Survey-weighted logistic regression results for core indicators.** Model 3 was the main adjusted model. Strict Model 4 added cardiometabolic, renal and inflammatory covariates and was interpreted as a conservative analysis because of complete-case reduction and potential overadjustment among correlated biomarkers.

**Table 5. GBD 2023 cardiovascular disease burden summary.** Age-standardised rates are shown for 1990 and 2023. Estimated annual percentage change (EAPC) was estimated from log-linear trends over calendar year.

**Table 6. Temporal external validation in NHANES 2021-2023.** Models were trained using NHANES 2005-2018 only and evaluated in NHANES 2021-2023. External weighted regression results are shown for UACR, eGFR, TyG-WC and SIRI.

## Figures

**Figure 1. Population-level cardiovascular disease burden and attributable risk context in GBD 2023.** (a) Age-standardised DALY, death and prevalence rates from 1990 to 2023 in Global, China and United States settings. (b) Leading all-age attributable CVD DALYs in 2023 for China and Global settings. DALY, disability-adjusted life-year.

**Figure 2. Internal held-out performance of the machine-learning models.** Receiver-operating characteristic curves, precision-recall curves and calibration plots are shown for the highest-performing full clinical, routine indicator and biomarker-focused models in the NHANES 2005-2018 held-out test set.

**Figure 3. SHAP explanations and consensus feature support for the biomarker-focused model.** (a) SHAP summary plot for the biomarker-focused model. (b) Consensus top-10 support across complementary feature-importance methods, with core indicators highlighted.

**Figure 4. Restricted cubic spline analyses for key indicators.** Survey-weighted natural spline curves are shown for UACR, eGFR, TyG-WC and SIRI using Model 3 adjustment. Curves represent odds ratios with shaded confidence intervals.

**Figure 5. Subgroup and sensitivity analyses for core indicators.** (a) Survey-weighted subgroup odds ratios for UACR, eGFR, TyG-WC and SIRI. (b) Sensitivity-analysis results after excluding prior stroke, excluding prior heart attack and using routine biomarker adjustment.

**Figure 6. Temporal external validation in NHANES 2021-2023.** External receiver-operating characteristic curves, precision-recall curves and calibration plots are shown for models trained only in NHANES 2005-2018 and tested in NHANES 2021-2023.
"""


def build_manuscript(references: str) -> str:
    return f"""# Renal, metabolic and inflammatory markers of cardiovascular disease

## Abstract

### Background

Cardiovascular disease (CVD) remains a major source of mortality and disability. Routine renal, metabolic and inflammatory indicators may capture complementary information relevant to CVD classification, but their relative value in nationally representative data remains uncertain.

### Methods

This public-database study used analytically separate population-level and individual-level data sources. GBD 2023 data described CVD burden trends and attributable risk-factor context for Global, China and United States settings. NHANES 2005-2018 adult data were used for leakage-audited machine learning, interpretable feature ranking and survey-weighted association analyses. NHANES 2021-2023 was reserved as a temporal external validation cohort. CVD was defined as self-reported congestive heart failure, coronary heart disease, angina, heart attack or myocardial infarction, or stroke. Three predefined feature sets were evaluated: a full clinical model, a routine indicator model and a biomarker-focused model. Frozen final models trained in NHANES 2005-2018 were tested in NHANES 2021-2023.

### Results

GBD 2023 showed declines in age-standardised CVD mortality and DALY rates between 1990 and 2023, although the public-health burden remained substantial. In NHANES 2005-2018, the analytic sample included 16,827 adults, including 1,942 with self-reported CVD. Held-out discrimination was similar across the highest-performing full clinical model (XGBoost; AUC 0.853, PR-AUC 0.442), routine indicator model (XGBoost; AUC 0.848, PR-AUC 0.439) and biomarker-focused model (Elastic Net; AUC 0.846, PR-AUC 0.418). Temporal external validation in NHANES 2021-2023 yielded AUCs of 0.845, 0.839 and 0.828, respectively. UACR showed the most consistent evidence across machine-learning prioritisation, survey-weighted regression, nonlinear analyses and sensitivity analyses. In external weighted regression, UACR (OR 1.001, 95% CI 1.000-1.001; P < 0.001), TyG-WC (OR 1.003, 95% CI 1.001-1.004; P = 0.0015) and SIRI (OR 1.285, 95% CI 1.084-1.522; P = 0.0039) were positively associated with CVD, whereas eGFR was inversely associated (OR 0.988, 95% CI 0.981-0.995; P < 0.001).

### Conclusions

GBD data provided population-level CVD context, whereas NHANES supported individual-level prediction, association analyses and temporal validation. A biomarker-focused model achieved discrimination close to broader clinical models, and UACR was the most robust cross-sectional indicator. eGFR, TyG-WC and SIRI were repeatedly prioritised but were more sensitive to adjustment and sensitivity analyses. These findings should be interpreted as hypothesis-generating evidence from cross-sectional public data, not as causal or implementation-ready clinical evidence.

### Keywords

Cardiovascular disease; NHANES; GBD; interpretable machine learning; UACR; eGFR; TyG-WC; SIRI; SHAP; survey-weighted analysis.

## Introduction

Cardiovascular disease remains a leading contributor to global mortality and disability. Population surveillance systems, including the Global Burden of Disease (GBD) study, provide a structured view of long-term changes in cardiovascular mortality, disability-adjusted life-years (DALYs), prevalence and risk-attributable burden [1,2]. These estimates are essential for public-health planning, but they do not directly identify individual-level indicators associated with CVD status.

Risk assessment for CVD relies on established demographic and clinical variables, including age, sex, blood pressure, diabetes, smoking and lipid measures. Many population studies also contain routinely measured renal, metabolic and inflammatory indicators that may capture related but non-identical biological information. The difficulty is not the absence of candidate markers, but the need to identify which markers remain repeatedly prioritised across modelling, regression and sensitivity analyses without overstating their independent clinical value.

Renal, metabolic and inflammatory pathways are closely connected with cardiovascular health [3-7]. Albuminuria, captured by the urine albumin-to-creatinine ratio (UACR), may reflect renal, endothelial and microvascular injury. Estimated glomerular filtration rate (eGFR) provides a complementary measure of kidney function [3,4]. Metabolic indices such as the triglyceride-glucose index and TyG-WC combine glycaemic, lipid and adiposity-related information [5,6]. Inflammatory indices such as the systemic inflammation response index (SIRI) summarise leukocyte patterns that may track chronic inflammatory states [7].

Interpretable machine learning can compare broad candidate-predictor sets while preserving feature-prioritisation tools such as permutation importance and SHAP-based explanations [8-10]. In the present study, GBD and NHANES were intentionally kept at different analytic levels. GBD was used only to describe population-level CVD burden and attributable risk context. NHANES was used for individual-level prediction, interpretable feature ranking and survey-weighted epidemiologic analyses. This separation avoids implying that macro-level burden estimates validate individual-level biomarkers.

This study integrated GBD 2023 CVD burden trends with NHANES 2005-2018 individual-level modelling to identify renal, metabolic and inflammatory indicators associated with self-reported CVD. We compared leakage-audited feature sets, summarised consensus feature rankings, evaluated survey-weighted associations, assessed nonlinear patterns, examined subgroup and sensitivity results, and tested temporal transportability in NHANES 2021-2023. The aim was to generate a cautious public-database evidence package for CVD risk-marker research rather than causal proof or a deployable clinical decision model.

## Methods

### Study design and data sources

This study had two analytically separate components. GBD 2023 CSV exports were used to describe CVD burden trends and attributable risk-factor rankings at the population level [1,2]. NHANES 2005-2018 data were used for individual-level machine learning, interpretable feature ranking and survey-weighted statistical analyses [11,12]. The two sources were not merged, and no joint GBD-NHANES model was fitted.

The separation of data sources preserved the intended level of inference. GBD files were treated as aggregate burden estimates and summarised by location, measure, metric, age category, sex and year. NHANES records were treated as individual-level survey observations and analysed with attention to sample weights, strata and primary sampling units. GBD findings were therefore used as population context, whereas NHANES analyses formed the basis for prediction and association results.

### GBD 2023 burden analysis

GBD 2023 data were obtained from official user-provided CSV exports through IHME/GBD tools [1,2]. The analysis focused on cardiovascular diseases and included deaths, DALYs and prevalence, with number and rate metrics. Locations included Global, China and United States. Age-standardised rates were summarised for 1990 and 2023, and percent change was computed directly from those values. Estimated annual percentage change (EAPC) was estimated from log-linear trends over calendar year. Attributable CVD burden was summarised for selected risk factors, including high systolic blood pressure, dietary risks, high LDL cholesterol, smoking, high fasting plasma glucose, kidney dysfunction and high body-mass index.

### NHANES study population

NHANES 2005-2018 adult participants aged 20 years or older were included according to the project pipeline. The primary analytic dataset used the fasting-core sample because TyG-related indicators required fasting glucose and triglycerides. Participants with missing CVD outcome were excluded during data preparation. The final analytic sample contained 16,827 adults. Survey-weighted analyses used the combined fasting weight, strata and primary sampling unit variables, consistent with NHANES guidance for combining cycles and selecting subsample weights [11,12].

NHANES 2021-2023 was prepared as an independent temporal external validation cohort and was not merged with NHANES 2005-2018 for training. The same self-reported CVD definition was applied. The external cohort used the fasting-core analytic subset when TyG-related indicators were required, and survey design variables were retained for external weighted regression. Final 2005-2018 models were refitted on the full training dataset and frozen before prediction in NHANES 2021-2023.

### CVD outcome definition

CVD was defined using self-reported physician diagnosis or history variables from the NHANES medical conditions questionnaire [14]. Participants were classified as having CVD if they reported congestive heart failure, coronary heart disease, angina, heart attack or myocardial infarction, or stroke. Outcome component variables and direct CVD diagnostic variables were excluded from predictor sets before model fitting.

### Candidate predictors and feature sets

Candidate predictors covered demographic, lifestyle, anthropometric, blood pressure, glucose and lipid metabolism, inflammatory and renal-function domains. Derived indicators included TyG index, TyG-BMI, TyG-WC, atherogenic index of plasma, non-HDL-C, neutrophil-to-lymphocyte ratio, platelet-to-lymphocyte ratio, systemic immune-inflammation index, SIRI, eGFR and UACR when source variables were available [3-7].

Three leakage-audited feature sets were evaluated. The full clinical model included demographic, lifestyle, anthropometric, blood pressure, metabolic, inflammatory, renal-function and selected comorbidity variables. The routine indicator model excluded hypertension and diabetes diagnosis variables and focused on routinely measured indicators. The biomarker-focused model included basic adjustment variables plus metabolic, inflammatory and renal indicators, including TyG-related indices, non-HDL-C, NLR, PLR, SII, SIRI, UACR, eGFR, uric acid, creatinine and albumin where available.

### Machine learning and interpretability

Machine-learning analyses compared logistic regression, random forest, XGBoost, LightGBM, support vector machine and elastic net classifiers across the three feature sets [8,15]. Data were split into stratified training and held-out test sets using a 70:30 ratio. Preprocessing was fitted within training-set pipelines only. Numeric variables were imputed with training-set medians; categorical variables were imputed with the training-set most frequent category and one-hot encoded; and standardisation was applied to models requiring scaled inputs. Class imbalance was handled through class weights or model-specific positive-class weighting. Model performance metrics included AUC, PR-AUC, F1 score, Brier score, sensitivity and specificity.

Feature importance was summarised using model-based importance, permutation importance and SHAP explanations computed on the held-out test set using models trained on the training set [9,10]. Consensus feature ranking was derived across methods within each feature set. No top-ranked feature subset was selected before splitting or used to refit the evaluated models.

### Survey-weighted statistical analysis

Survey-weighted logistic regression estimated odds ratios (ORs) and 95% confidence intervals (CIs) for key indicators using complex-survey methods [16]. Model 3 was the main adjusted model and included core demographic and lifestyle factors. Strict Model 4 added blood pressure, glycaemic, lipid, renal and inflammatory covariates while excluding the target indicator itself. Model 4 was interpreted as a conservative analysis because it used a smaller complete-case sample and may overadjust among correlated biomarkers.

Nonlinear associations were assessed using survey-weighted natural spline models for prespecified indicators, including UACR, eGFR, TyG-WC, SIRI, TyG index and non-HDL-C [17]. Subgroup analyses evaluated age group, sex, diabetes status, hypertension status and obesity status. Sensitivity analyses excluded prior stroke, excluded prior heart attack and used routine biomarker adjustment without hypertension or diabetes diagnosis variables.

### Ethics and software

NHANES protocols were approved by the National Center for Health Statistics ethics review board, and participants provided informed consent at the time of data collection. The present secondary analysis used de-identified public NHANES data and official GBD CSV exports. Additional institutional review board approval was not required.

Reporting was checked against TRIPOD+AI for prediction modelling and STROBE for observational epidemiology [18,19]. Data preparation, machine learning, feature ranking, GBD analysis, figure generation and manuscript table generation were performed using Python. Survey-weighted analyses were performed using R with the `survey` package, and natural spline analyses used the R `splines` package.

## Results

### GBD burden trends

GBD 2023 data showed declining age-standardised CVD death and DALY rates between 1990 and 2023 in Global, China and United States settings (Table 5, Figure 1). Global CVD death rates decreased from 375.51 to 214.94 per 100,000 (-42.8%), and global DALY rates decreased from 8061.62 to 4863.82 per 100,000 (-39.7%). In China, age-standardised death rates decreased from 496.65 to 208.03 per 100,000 (-58.1%), and DALY rates decreased from 10065.67 to 4059.34 per 100,000 (-59.7%). China showed a small increase in age-standardised prevalence rate, from 6046.60 to 6180.58 per 100,000 (2.2%), with a stable EAPC estimate.

The leading attributable CVD DALY risk factors in 2023 were broadly similar across Global and China settings. In China, the leading risk factors by all-age DALY number were high systolic blood pressure, dietary risks, high LDL cholesterol, smoking and high fasting plasma glucose. Globally, the leading risk factors were high systolic blood pressure, dietary risks, high LDL cholesterol, smoking and high body-mass index. These GBD findings were used only to frame the population-level burden context.

### NHANES population characteristics

The NHANES analytic sample included 16,827 adults, of whom 1,942 had self-reported CVD and 14,885 did not (Table 1). Participants with CVD were older than those without CVD (64.31 versus 45.82 years, weighted means) and had higher body mass index, larger waist circumference, higher systolic blood pressure, higher fasting glucose and higher HbA1c. Renal and inflammatory indicators also differed between groups, including lower eGFR, higher UACR, higher serum creatinine, higher NLR and higher SIRI among participants with CVD.

### Machine-learning model performance

The leakage-audited machine-learning models showed similar held-out discrimination across the three feature sets (Table 2, Figure 2). In the full clinical feature set, XGBoost achieved AUC 0.853, PR-AUC 0.442, F1 0.444 and Brier score 0.144. In the routine indicator feature set, XGBoost achieved AUC 0.848, PR-AUC 0.439, F1 0.444 and Brier score 0.146. In the biomarker-focused feature set, Elastic Net achieved AUC 0.846, PR-AUC 0.418, F1 0.421 and Brier score 0.166. The AUC difference between the full clinical and biomarker-focused models was 0.007.

### Temporal external validation

NHANES 2021-2023 external validation included 3,055 adults, including 384 participants with CVD (Table 6, Figure 6). The final full clinical XGBoost model achieved external AUC 0.845, PR-AUC 0.468, F1 0.445 and Brier score 0.169. The routine indicator XGBoost model achieved external AUC 0.839, PR-AUC 0.452, F1 0.438 and Brier score 0.172. The biomarker-focused Elastic Net model achieved external AUC 0.828, PR-AUC 0.402, F1 0.409 and Brier score 0.193. External minus internal AUC differences were small for the full clinical, routine indicator and biomarker-focused models (-0.007, -0.009 and -0.018, respectively).

In external weighted regression, UACR, TyG-WC and SIRI showed positive associations with CVD, whereas eGFR showed an inverse association. These analyses used NHANES 2021-2023 only as a temporal validation cohort and should not be interpreted as prospective outcome validation.

### Feature importance and statistical validation

Consensus feature ranking highlighted renal, metabolic, inflammatory and lipid-related indicators (Table 3, Figure 3). UACR entered the top 10 in both routine and biomarker-focused models. eGFR was ranked second in both feature sets. TyG-WC also entered the top 10 in both models. SIRI was not in the routine top 10 but was ranked seventh in the biomarker-focused model. TyG index and non-HDL-C were prioritised in the biomarker-focused top 10, whereas LDL-C and HbA1c were prioritised in the routine top 10.

In survey-weighted logistic regression, UACR showed the most robust evidence across models (Table 4). In Model 3, UACR was associated with CVD (OR 1.001, 95% CI 1.001-1.002; P < 0.001), and the association remained present in strict Model 4 (OR 1.001, 95% CI 1.000-1.002; P = 0.0219). eGFR was associated with CVD in Model 3 (OR 0.983, 95% CI 0.978-0.988; P < 0.001) but was attenuated in strict Model 4. TyG-WC and SIRI also showed Model 3 associations but did not remain statistically significant in strict Model 4.

Restricted cubic spline analyses supported nonlinear association patterns for UACR, eGFR, TyG-WC and SIRI (Figure 4). Subgroup analysis identified an interaction for TyG-WC by obesity status. In sensitivity analyses, UACR remained statistically significant in 2 of 3 sensitivity analyses, whereas eGFR was significant in 0 of 3, TyG-WC in 1 of 3 and SIRI in 1 of 3 (Figure 5). These findings supported UACR as the most stable indicator across analytic layers, while eGFR, TyG-WC and SIRI were supported but more sensitive to model specification.

## Discussion

This study combined population-level GBD context with individual-level NHANES modelling while maintaining a strict separation between the two levels of inference. GBD 2023 showed declining age-standardised CVD death and DALY rates from 1990 to 2023, with high systolic blood pressure, dietary risks, high LDL cholesterol and smoking remaining leading attributable risk factors. In NHANES 2005-2018, interpretable machine-learning models showed comparable discrimination across full clinical, routine indicator and biomarker-focused feature sets. UACR emerged as the most robust individual-level indicator, whereas eGFR, TyG-WC and SIRI were repeatedly prioritised but less stable under strict adjustment or sensitivity analyses.

The comparable discrimination of the full clinical and biomarker-focused models suggests that a focused set of routine renal, metabolic and inflammatory indicators captured substantial information relevant to CVD classification in this cross-sectional NHANES setting. This does not mean that a biomarker-focused model can replace comprehensive clinical assessment. Rather, it suggests that routinely available biomarkers may be useful for risk-marker screening, phenotyping and hypothesis generation. Clinical utility would require calibration assessment, threshold selection, prospective validation and evaluation against established risk scores.

UACR showed the most consistent evidence across analytic layers. It was prioritised by machine learning, remained associated in the strict fully adjusted model, showed nonlinearity and retained partial support in sensitivity analyses. This pattern is consistent with the concept that albuminuria may capture renal, endothelial and microvascular information not fully summarised by standard covariates [4]. The present study does not establish that UACR causes future cardiovascular events or should be used alone for clinical decisions. It supports UACR as a robust cross-sectional marker associated with self-reported CVD in a nationally representative adult sample.

eGFR, TyG-WC and SIRI should be interpreted as complementary signals rather than fully independent clinical predictors. eGFR provided a strong renal-function signal in machine-learning ranking and Model 3 regression, but its estimate was attenuated in strict Model 4. TyG-WC integrated metabolic and anthropometric information and was supported by feature ranking, Model 3 regression, nonlinear analysis and an obesity-status interaction. SIRI represented the inflammatory domain and was supported by biomarker-focused ranking, Model 3 regression and nonlinear analysis. Their attenuation under extensive mutual adjustment highlights biological and statistical overlap among routine biomarkers.

The GBD results provide population-level context rather than validation of NHANES biomarkers [1,2]. GBD attributable burden describes macro-level patterns and prevention priorities. NHANES evaluates associations between individual-level indicators and self-reported CVD status. Placing these results side by side is useful because it frames individual-level marker analysis within the broader public-health burden, but the two evidence streams should not be merged into a single causal interpretation.

Temporal validation in NHANES 2021-2023 supported cautious transportability across NHANES survey periods. External AUC declines were small for all three final models, and the four key indicators showed directions consistent with the main analysis. The validation nevertheless remains cross-sectional and based on a later NHANES survey cycle rather than adjudicated incident cardiovascular events. CRP and older physical activity variables were not fully comparable in NHANES 2021-2023, and survey-period differences may affect model transportability.

This study has several strengths. It used large, public and reproducible data sources; kept GBD and NHANES analytically separate; incorporated leakage auditing; evaluated prespecified feature sets; used held-out and temporal validation; and triangulated findings across machine learning, regression, nonlinear, subgroup and sensitivity analyses. The survey-weighted analyses respected the complex NHANES sampling design.

Several limitations should be considered. NHANES CVD status was self-reported and cross-sectional, so temporality and causality cannot be established. External validation used a later NHANES cycle, but it was not prospective event validation and remained within a US survey system. Medication use, treatment history and clinical management may influence biomarker levels, especially lipid measures. Strict Model 4 had a smaller complete-case sample and may overadjust correlated biomarkers. Some indicators were unavailable for all participants, and missingness may affect complete-case survey models. RCS analyses were limited to prespecified indicators. Non-US cohorts and prospective datasets with adjudicated cardiovascular outcomes are needed before clinical implementation can be considered.

Future work should test whether UACR, eGFR, TyG-WC and SIRI improve prediction beyond established risk factors in longitudinal cohorts. Additional studies should evaluate medication use, kidney disease stage, diabetes subgroups, race and ethnicity, recalibration across populations and decision-analytic performance. Until such evidence is available, these indicators should be viewed as cross-sectional risk-marker signals rather than clinical decision rules.

## Conclusions

In public GBD and NHANES data, population-level CVD burden trends and individual-level biomarker associations provided complementary evidence. A biomarker-focused model achieved discrimination close to broader clinical models and retained reasonable temporal transportability. UACR was the most robust cross-sectional indicator, while eGFR, TyG-WC and SIRI were repeatedly prioritised but more sensitive to model specification. Prospective validation with adjudicated outcomes is required before clinical use.

## Tables and Figures

### Tables

- Table 1. Survey-weighted baseline characteristics according to CVD status.
- Table 2. Held-out machine-learning model performance by feature set.
- Table 3. Feature-prioritisation and statistical validation summary.
- Table 4. Survey-weighted regression results for core indicators.
- Table 5. GBD 2023 CVD burden summary.
- Table 6. NHANES 2021-2023 temporal external validation.

### Figures

- Figure 1. Population-level CVD burden and attributable risk context.
- Figure 2. Internal held-out machine-learning model performance.
- Figure 3. SHAP explanations and consensus feature support.
- Figure 4. Restricted cubic spline analyses for key indicators.
- Figure 5. Subgroup and sensitivity analyses.
- Figure 6. NHANES 2021-2023 external validation performance.

## Declarations

### Ethics approval and consent to participate

NHANES protocols were approved by the National Center for Health Statistics ethics review board, and participants provided informed consent at the time of data collection. This secondary analysis used de-identified public NHANES data and official GBD CSV exports.

### Consent for publication

Not applicable.

### Availability of data and materials

NHANES data are publicly available from the National Center for Health Statistics. GBD results are available through official IHME/GBD tools. Derived source tables underlying the figures are provided with this submission package. Analysis scripts and derived outputs should be deposited in a public repository before final publication, with the repository identifier added to this section.

### Competing interests

Author information has been removed for peer review.

### Funding

Author information has been removed for peer review.

### Authors' contributions

Author information has been removed for peer review.

### Acknowledgements

The authors acknowledge the participants and staff of NHANES and the IHME/GBD collaborators who produced the public data resources used in this study.

{references.strip()}
"""


def get_references() -> str:
    text = (SOURCE_PACKAGE / "blinded_manuscript.md").read_text(encoding="utf-8")
    marker = "## References"
    if marker not in text:
        raise ValueError("References marker not found in blinded_manuscript.md")
    return text[text.index(marker) :]


def create_docx_from_markdown(markdown_text: str, out_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    for raw in markdown_text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            add_inline_runs(p, line)
    doc.save(out_path)


def create_review_docx_with_figures(markdown_text: str, out_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    for raw in markdown_text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            add_inline_runs(p, line)

    doc.add_page_break()
    doc.add_heading("Embedded Figures for Review", level=1)
    figure_titles = [
        ("Figure 1", "Figure_1_GBD_CVD_burden_and_risk_factors.png"),
        ("Figure 2", "Figure_2_ML_model_performance.png"),
        ("Figure 3", "Figure_3_SHAP_and_feature_importance.png"),
        ("Figure 4", "Figure_4_RCS_key_indicators.png"),
        ("Figure 5", "Figure_5_Subgroup_and_sensitivity_analysis.png"),
        ("Figure 6", "Figure_6_External_validation_performance.png"),
    ]
    for title, filename in figure_titles:
        path = DEST_FIGURES / filename
        if not path.exists():
            continue
        doc.add_heading(title, level=2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Inches(6.5))
    doc.save(out_path)


def add_inline_runs(paragraph, line: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", line)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def create_tables_docx(table_outputs: dict[str, pd.DataFrame]) -> None:
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.55)
    sec.right_margin = Inches(0.55)
    doc.add_heading("Publication Tables", level=0)
    for name, df in table_outputs.items():
        doc.add_heading(name, level=1)
        table = doc.add_table(rows=1, cols=len(df.columns))
        table.style = "Table Grid"
        for j, col in enumerate(df.columns):
            table.rows[0].cells[j].text = str(col)
        for _, row in df.iterrows():
            cells = table.add_row().cells
            for j, value in enumerate(row):
                cells[j].text = "" if pd.isna(value) else str(value)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(8)
        doc.add_paragraph()
    doc.save(DEST_TABLES / "publication_tables_polished.docx")


def create_manuscript_files() -> None:
    references = get_references()
    manuscript = build_manuscript(references)
    legends = build_legends()
    (DEST_PACKAGE / "blinded_manuscript_nature_polished.md").write_text(manuscript, encoding="utf-8")
    with_figures_md = manuscript + "\n\n## Embedded Figures for Review\n\n"
    for idx in range(1, 7):
        match = next(DEST_FIGURES.glob(f"Figure_{idx}_*.png"), None)
        if match is not None:
            with_figures_md += f"### Figure {idx}\n\n![Figure {idx}](figures/{match.name})\n\n"
    (DEST_PACKAGE / "blinded_manuscript_nature_polished_with_figures.md").write_text(with_figures_md, encoding="utf-8")
    (DEST_PACKAGE / "table_figure_legends_nature_polished.md").write_text(legends, encoding="utf-8")
    create_docx_from_markdown(manuscript, DEST_PACKAGE / "blinded_manuscript_nature_polished.docx")
    create_review_docx_with_figures(manuscript, DEST_PACKAGE / "blinded_manuscript_nature_polished_with_figures.docx")
    create_docx_from_markdown(legends, DEST_PACKAGE / "table_figure_legends_nature_polished.docx")


def create_revision_report(table_outputs: dict[str, pd.DataFrame]) -> None:
    figure_files = sorted(DEST_FIGURES.glob("Figure_*.png"))
    pdf_files = sorted(DEST_FIGURES.glob("Figure_*.pdf"))
    svg_files = sorted(DEST_FIGURES.glob("Figure_*.svg"))
    report = f"""# Nature-Style Revision Report

## Scope

- Rewrote the manuscript for a clearer SCI/Nature-leaning argument structure: context, gap, approach, evidence, interpretation and boundary.
- Softened causal and clinical implementation claims, preserving the cross-sectional and public-database nature of the evidence.
- Reformatted tables for publication readability and corrected lipid-unit labels in Table 1 from mmol/L to mg/dL, matching the reported NHANES concentration scale.
- Rebuilt figures with a consistent Python/matplotlib backend, restrained palette, smaller typography, reduced whitespace and 600 dpi raster exports.
- Added source-data copies for figure regeneration and audit.

## Generated files

- Polished manuscript: `blinded_manuscript_nature_polished.docx` and `.md`
- Polished legends: `table_figure_legends_nature_polished.docx` and `.md`
- Tables: {len(table_outputs)} polished CSV files, one styled Excel workbook and one Word table file
- Figures: {len(figure_files)} high-resolution PNG files, {len(pdf_files)} PDF files and {len(svg_files)} SVG files

## Remaining submission checks

- Add a public repository identifier or DOI for analysis scripts and derived outputs before final publication.
- Confirm the target journal's preferred abstract structure, word limit and reference style.
- Confirm whether figures should be submitted as TIFF only or as editable PDF/SVG plus source data.
- Have a human author verify every numeric statement against the final locked analysis outputs before submission.
"""
    (DEST_PACKAGE / "nature_polishing_revision_report.md").write_text(report, encoding="utf-8")


def main() -> None:
    ensure_dirs()
    configure_matplotlib()
    table_outputs = polish_tables()
    create_excel_workbook(table_outputs)
    create_tables_docx(table_outputs)
    copy_source_data()
    create_figures()
    create_manuscript_files()
    create_revision_report(table_outputs)
    print(f"Polished package written to: {DEST_PACKAGE}")


if __name__ == "__main__":
    main()
